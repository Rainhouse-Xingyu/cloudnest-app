#!/usr/bin/env python3
"""Generate public synthetic OOXML fixtures; requires python-docx/openpyxl/python-pptx.
No model responses are stored. Fixed ZIP metadata makes output reviewable.
"""
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZipInfo, ZIP_STORED, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

out = Path(__file__).resolve().parents[1] / 'cloudnest_go/integration/testdata/office'
out.mkdir(parents=True, exist_ok=True)
fixed = datetime(2020, 1, 1)
with TemporaryDirectory() as directory:
    root = Path(directory)
    doc = Document()
    doc.core_properties.author = 'CloudNest synthetic fixtures'
    doc.core_properties.created = doc.core_properties.modified = fixed
    doc.add_heading('数据库迁移方案', 0)
    p = doc.add_paragraph('MySQL 升级前先')
    p.add_run('备份数据库').bold = True
    p.add_run('，迁移后验证数据一致性。')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0,0).text = '服务器 IP'
    table.cell(0,1).text = '192.0.2.42'
    table.cell(1,0).text = '回滚计划'
    table.cell(1,1).text = '恢复备份，核对行数和校验值。'
    doc.sections[0].header.paragraphs[0].text = 'CloudNest 合成测试文档'
    doc.save(root / 'migration.docx')

    book = Workbook()
    book.properties.creator = 'CloudNest synthetic fixtures'
    book.properties.created = book.properties.modified = fixed
    sheet = book.active
    sheet.title = '电脑购买发票'
    sheet.append(['商品', '金额', '测试说明'])
    sheet.append(['笔记本电脑', 5999, '合成票据，无个人数据'])
    sheet.append(['合计', '=SUM(B2:B2)', '服务器 IP 192.0.2.42'])
    book.save(root / 'invoice.xlsx')

    deck = Presentation()
    deck.core_properties.author = 'CloudNest synthetic fixtures'
    deck.core_properties.created = deck.core_properties.modified = fixed
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = 'Spring Boot 权限管理：角色、资源授权与审计日志。'
    slide.notes_slide.notes_text_frame.text = '演讲备注：检查文件级访问权限，阻止跨用户读取。'
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = '安全回归测试与部署检查。'
    deck.save(root / 'permissions.pptx')

    for source in sorted(root.iterdir()):
        with ZipFile(source) as src, ZipFile(out / source.name, 'w') as dst:
            for name in sorted(src.namelist()):
                info = ZipInfo(name, (2020,1,1,0,0,0))
                # Valid uncompressed DOCX deliberately exceeds the text-only
                # 256 KiB source limit while its extracted text is small.
                info.compress_type = ZIP_STORED if source.suffix == '.docx' else ZIP_DEFLATED
                data = src.read(name)
                if name == 'docProps/core.xml':
                    tree = ET.fromstring(data)
                    for tag in ('created', 'modified'):
                        node = tree.find('{http://purl.org/dc/terms/}' + tag)
                        if node is not None:
                            node.text = '2020-01-01T00:00:00Z'
                    data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
                dst.writestr(info, data)
        print(source.name, (out / source.name).stat().st_size)
